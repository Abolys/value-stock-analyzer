"""Word export (python-docx). The footer sits in every section's page footer, so
Word prints it on every page."""

from __future__ import annotations

import io

from docx import Document
from docx.shared import Inches, Pt

from reports.build import Report, plain

PICTURE_WIDTH = Inches(6.3)


def _set_footer(doc, text: str) -> None:
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        para.text = text
        for run in para.runs:
            run.font.size = Pt(8)


def to_docx(rep: Report) -> bytes:
    doc = Document()
    doc.add_heading(rep.title, level=0)
    for sec in rep.sections:
        doc.add_heading(sec.title, level=1)
        for p in sec.paragraphs:
            for chunk in plain(p).split("\n\n"):
                if chunk.strip():
                    doc.add_paragraph(chunk.strip())
        for t in sec.tables:
            table = doc.add_table(rows=1, cols=len(t.headers))
            table.style = "Table Grid"
            for cell, h in zip(table.rows[0].cells, t.headers):
                cell.text = h
            for row in t.rows:
                cells = table.add_row().cells
                for cell, v in zip(cells, row):
                    cell.text = plain(str(v))
            doc.add_paragraph()
        for f in sec.figures:
            if f.png is not None:
                doc.add_picture(io.BytesIO(f.png), width=PICTURE_WIDTH)
            else:
                doc.add_paragraph(f"{f.title} chart not rendered: {f.reason}")
            if f.caption:
                cap = doc.add_paragraph(f.caption)
                for run in cap.runs:
                    run.font.size = Pt(8)
    _set_footer(doc, rep.footer)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
