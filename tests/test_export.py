"""Per-ticker export: the .docx opens and contains every section, the footer is on
every page (every docx section footer; after every Markdown section), assumptions
and data gaps are included, and charts that can't be rendered are listed."""

import io
from functools import lru_cache
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from app import stock_view as sv
from data.fixture_provider import captured_on, fixture_edgar, fixture_provider
from reports.build import DISCLAIMER, build_report
from reports.docx_export import to_docx
from reports.markdown import to_markdown

SECTIONS = ["Summary", "Lens scores", "Quant fundamental lens", "Macro & balance sheet lens", "Business moat lens",
            "Devil's advocate lens", "Fundamentals over time", "Valuation", "Versus peers", "Asset floor",
            "Value-trap scores, EV/EBIT and insiders", "Dividend and context", "Turnaround outlook", "Assumptions",
            "Data gaps"]


def tiny_png(_fig):
    buf = io.BytesIO()
    Image.new("RGB", (4, 3), "white").save(buf, format="PNG")
    return buf.getvalue(), ""


def broken(_fig):
    return None, "RuntimeError: Chrome not found"


@lru_cache(maxsize=None)
def _lulu(tmp: str):
    from analysis.pipeline import run_analysis
    from llm.cache import LLMCache
    from llm.client import LLMClient
    from screening.engine import ScreenContext
    from tests.llm_fakes import FakeAPI

    db = Path(tmp) / "runs.db"
    provider = fixture_provider()
    llm = LLMClient(api=FakeAPI(), cache=LLMCache(Path(tmp) / "llm.db"), db_path=db)
    run = run_analysis(ScreenContext(provider=provider, db_path=db, today=captured_on("LULU")), "LULU", llm=llm,
                       edgar=fixture_edgar())
    bundle = sv.load_bundle(provider, run, db)
    return run, sv.build_charts(run, bundle)


@pytest.fixture
def lulu(tmp_path_factory):
    return _lulu(str(tmp_path_factory.getbasetemp() / "export"))


def test_docx_opens_and_contains_every_section_with_footer_on_every_page(lulu):
    run, charts = lulu
    rep = build_report(run, charts, render=tiny_png)
    doc = Document(io.BytesIO(to_docx(rep)))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for s in SECTIONS:
        assert s in headings, s
    assert doc.sections, "no sections"
    for section in doc.sections:
        text = " ".join(p.text for p in section.footer.paragraphs)
        assert DISCLAIMER in text and "Price as of" in text and "Fundamentals as of" in text
        assert "Value Stock Analyzer v" in text
    assert len(doc.inline_shapes) >= 5  # charts embedded as pictures
    body = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "COST_OF_CAPITAL" in cells  # assumptions table
    assert "How this score was built" in body
    assert run.turnaround.survivorship_caveat in body


def test_markdown_footer_after_every_section(lulu):
    run, charts = lulu
    rep = build_report(run, charts, render=tiny_png)
    md = to_markdown(rep)
    assert md.count(DISCLAIMER) == len(rep.sections)
    assert md.rstrip().endswith(f"_{rep.footer}_")
    for s in SECTIONS:
        assert f"## {s}" in md
    assert "data:image/png;base64," in md


def test_data_gaps_and_unrendered_charts_listed(lulu):
    run, charts = lulu
    rep = build_report(run, charts, render=broken)
    gaps = next(s for s in rep.sections if s.title == "Data gaps")
    rows = {tuple(r) for t in gaps.tables for r in t.rows}
    assert ("Business Moat", run.moat.status) in rows or run.moat.ok  # an LLM lens gap is listed if any
    assert any("not rendered: RuntimeError: Chrome not found" in r[1] for r in rows)
    # LULU files with the SEC: its valuation-based recovery is computed from XBRL history, so it's not a gap
    assert run.turnaround.valuation.status in ("ok", "not_cheap")
    assert not any(r[0] == "Valuation-based recovery" for r in rows)
    doc = Document(io.BytesIO(to_docx(rep)))
    assert any("chart not rendered" in p.text for p in doc.paragraphs)


def test_footer_names_dates_providers_and_version(lulu):
    import config

    run, _ = lulu
    rep = build_report(run, {}, render=tiny_png)
    assert f"Price as of {run.price_as_of}" in rep.footer
    assert f"Fundamentals as of {run.fundamentals_as_of}" in rep.footer
    assert all(p in rep.footer for p in run.providers) and f"v{config.APP_VERSION}" in rep.footer


def test_project_chrome_is_used_unless_browser_path_is_set(tmp_path, monkeypatch):
    import os

    import config
    from reports import images

    exe = tmp_path / "chrome-linux64" / "chrome"
    exe.parent.mkdir()
    exe.write_text("")
    monkeypatch.setattr(config, "CHROME_DIR", tmp_path)
    saved = os.environ.pop("BROWSER_PATH", None)  # use_project_chrome sets it directly: restore by hand
    try:
        assert images.project_chrome() == exe
        assert images.use_project_chrome() == str(exe) and os.environ["BROWSER_PATH"] == str(exe)
        os.environ["BROWSER_PATH"] = "/opt/my/chrome"
        assert images.use_project_chrome() == "/opt/my/chrome"  # an explicit override wins
        del os.environ["BROWSER_PATH"]
        monkeypatch.setattr(config, "CHROME_DIR", tmp_path / "empty")
        assert images.project_chrome() is None and images.use_project_chrome() == ""
    finally:
        os.environ.pop("BROWSER_PATH", None)
        if saved is not None:
            os.environ["BROWSER_PATH"] = saved
